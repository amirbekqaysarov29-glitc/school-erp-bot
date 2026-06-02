import asyncio
import os
import random
import sqlite3
import shutil
import threading
import html
from pathlib import Path
from datetime import date, datetime, timedelta

import openpyxl
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart
from aiogram.types import Message, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton, FSInputFile

from fastapi import FastAPI
from fastapi.responses import HTMLResponse
import uvicorn

try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass


# ================= CONFIG =================

BOT_TOKEN = os.getenv("BOT_TOKEN", "").strip()
ADMIN_CODE = os.getenv("ADMIN_CODE", "ADMIN2026")
TEACHER_CODE = os.getenv("TEACHER_CODE", "TEACHER2026")

DB_NAME = os.getenv("DB_NAME", "data/davomat.db")

# Local: http://127.0.0.1:8000
# Deploy: set WEB_HOST=0.0.0.0
WEB_HOST = os.getenv("WEB_HOST", "0.0.0.0")
WEB_PORT = int(os.getenv("WEB_PORT", "8000"))

Path("data").mkdir(exist_ok=True)
Path("exports/excel").mkdir(parents=True, exist_ok=True)
Path("exports/pdf").mkdir(parents=True, exist_ok=True)
Path("backups").mkdir(exist_ok=True)

web_app = FastAPI(title="School ERP Dashboard")


# ================= DATABASE =================

def db():
    conn = sqlite3.connect(DB_NAME)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    c = conn.cursor()

    c.execute("""
    CREATE TABLE IF NOT EXISTS users (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        telegram_id INTEGER UNIQUE,
        role TEXT,
        full_name TEXT,
        created_at TEXT
    )
    """)

    c.execute("""
    CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        full_name TEXT,
        class_name TEXT,
        parent_phone TEXT,
        parent_telegram_id INTEGER,
        parent_code TEXT UNIQUE,
        created_at TEXT
    )
    """)

    c.execute("""
    CREATE TABLE IF NOT EXISTS teachers (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        full_name TEXT,
        telegram_id INTEGER UNIQUE,
        teacher_code TEXT UNIQUE,
        created_at TEXT
    )
    """)

    c.execute("""
    CREATE TABLE IF NOT EXISTS teacher_classes (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        teacher_id INTEGER,
        class_name TEXT
    )
    """)

    c.execute("""
    CREATE TABLE IF NOT EXISTS subjects (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE,
        created_at TEXT
    )
    """)

    c.execute("""
    CREATE TABLE IF NOT EXISTS attendance (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        student_id INTEGER,
        full_name TEXT,
        class_name TEXT,
        subject TEXT,
        date TEXT,
        status TEXT,
        teacher_telegram_id INTEGER,
        created_at TEXT
    )
    """)

    # Default subjects. Teacher can add more.
    for subject_name in ["Matematika", "Ona tili", "Ingliz tili"]:
        c.execute(
            "INSERT OR IGNORE INTO subjects(name, created_at) VALUES (?, ?)",
            (subject_name, datetime.now().isoformat())
        )

    conn.commit()
    conn.close()


def unique_code(cur, table, column, prefix):
    while True:
        code = f"{prefix}-{random.randint(1000, 9999)}"
        cur.execute(f"SELECT id FROM {table} WHERE {column} = ?", (code,))
        if not cur.fetchone():
            return code


def set_role(telegram_id, user_role, full_name=None):
    conn = db()
    c = conn.cursor()

    c.execute("""
    INSERT INTO users (telegram_id, role, full_name, created_at)
    VALUES (?, ?, ?, ?)
    ON CONFLICT(telegram_id)
    DO UPDATE SET
        role = excluded.role,
        full_name = COALESCE(excluded.full_name, users.full_name)
    """, (
        telegram_id,
        user_role,
        full_name,
        datetime.now().isoformat()
    ))

    conn.commit()
    conn.close()


def role(telegram_id):
    conn = db()
    c = conn.cursor()

    c.execute("SELECT role FROM users WHERE telegram_id = ?", (telegram_id,))
    row = c.fetchone()

    conn.close()
    return row["role"] if row else None


def ids_by_role(user_role):
    conn = db()
    c = conn.cursor()

    c.execute("SELECT telegram_id FROM users WHERE role = ?", (user_role,))
    rows = c.fetchall()

    conn.close()
    return [row["telegram_id"] for row in rows]


def add_student(full_name, class_name, phone):
    conn = db()
    c = conn.cursor()

    code = unique_code(c, "students", "parent_code", "PAR")

    c.execute("""
    INSERT INTO students (full_name, class_name, parent_phone, parent_code, created_at)
    VALUES (?, ?, ?, ?, ?)
    """, (
        full_name.strip(),
        class_name.strip(),
        phone.strip(),
        code,
        datetime.now().isoformat()
    ))

    conn.commit()
    conn.close()
    return code


def students():
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT *
    FROM students
    ORDER BY class_name, full_name
    """)

    rows = c.fetchall()
    conn.close()
    return rows


def student(student_id):
    conn = db()
    c = conn.cursor()

    c.execute("SELECT * FROM students WHERE id = ?", (student_id,))
    row = c.fetchone()

    conn.close()
    return row


def search_students(query):
    conn = db()
    c = conn.cursor()

    like = f"%{query.strip()}%"

    c.execute("""
    SELECT *
    FROM students
    WHERE full_name LIKE ?
       OR class_name LIKE ?
       OR parent_phone LIKE ?
       OR parent_code LIKE ?
    ORDER BY class_name, full_name
    """, (like, like, like, like))

    rows = c.fetchall()
    conn.close()
    return rows


def delete_student(student_id):
    conn = db()
    c = conn.cursor()

    c.execute("DELETE FROM students WHERE id = ?", (student_id,))
    count = c.rowcount

    conn.commit()
    conn.close()
    return count


def classes():
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT DISTINCT class_name
    FROM students
    ORDER BY class_name
    """)

    rows = c.fetchall()
    conn.close()
    return [row["class_name"] for row in rows]


def teacher_classes(telegram_id):
    conn = db()
    c = conn.cursor()

    c.execute("SELECT id FROM teachers WHERE telegram_id = ?", (telegram_id,))
    teacher = c.fetchone()

    if not teacher:
        conn.close()
        return []

    c.execute("""
    SELECT class_name
    FROM teacher_classes
    WHERE teacher_id = ?
    ORDER BY class_name
    """, (teacher["id"],))

    rows = c.fetchall()
    conn.close()

    return [row["class_name"] for row in rows]


def classes_for_teacher(telegram_id):
    allowed = teacher_classes(telegram_id)
    return allowed if allowed else classes()


def students_by_class(class_name):
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT id, full_name, class_name
    FROM students
    WHERE class_name = ?
    ORDER BY full_name
    """, (class_name,))

    rows = c.fetchall()
    conn.close()
    return rows


def link_parent(code, telegram_id, full_name=None):
    conn = db()
    c = conn.cursor()

    c.execute("""
    UPDATE students
    SET parent_telegram_id = ?
    WHERE UPPER(parent_code) = ?
    """, (
        telegram_id,
        code.strip().upper()
    ))

    count = c.rowcount
    conn.commit()
    conn.close()

    if count:
        set_role(telegram_id, "parent", full_name)

    return count


def parent_children(telegram_id):
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT *
    FROM students
    WHERE parent_telegram_id = ?
    ORDER BY class_name, full_name
    """, (telegram_id,))

    rows = c.fetchall()
    conn.close()
    return rows


def add_subject(name):
    conn = db()
    c = conn.cursor()

    clean_name = name.strip()
    if clean_name:
        c.execute("""
        INSERT OR IGNORE INTO subjects (name, created_at)
        VALUES (?, ?)
        """, (
            clean_name,
            datetime.now().isoformat()
        ))

    conn.commit()
    conn.close()


def subjects():
    conn = db()
    c = conn.cursor()

    c.execute("SELECT name FROM subjects ORDER BY name")
    rows = c.fetchall()

    conn.close()
    return [row["name"] for row in rows]


def add_teacher(name, classes_text=""):
    conn = db()
    c = conn.cursor()

    code = unique_code(c, "teachers", "teacher_code", "TCH")

    c.execute("""
    INSERT INTO teachers (full_name, teacher_code, created_at)
    VALUES (?, ?, ?)
    """, (
        name.strip(),
        code,
        datetime.now().isoformat()
    ))

    teacher_id = c.lastrowid

    for class_name in [item.strip() for item in classes_text.split(",") if item.strip()]:
        c.execute("""
        INSERT INTO teacher_classes (teacher_id, class_name)
        VALUES (?, ?)
        """, (
            teacher_id,
            class_name
        ))

    conn.commit()
    conn.close()
    return code


def link_teacher(code, telegram_id, full_name=None):
    conn = db()
    c = conn.cursor()

    c.execute("""
    UPDATE teachers
    SET telegram_id = ?
    WHERE UPPER(teacher_code) = ?
    """, (
        telegram_id,
        code.strip().upper()
    ))

    count = c.rowcount
    conn.commit()
    conn.close()

    if count:
        set_role(telegram_id, "teacher", full_name)

    return count


def teachers():
    conn = db()
    c = conn.cursor()

    c.execute("SELECT * FROM teachers ORDER BY full_name")
    rows = c.fetchall()

    conn.close()
    return rows


def save_attendance(student_id, full_name, class_name, subject_name, status, teacher_telegram_id):
    conn = db()
    c = conn.cursor()

    today = str(date.today())

    c.execute("""
    INSERT INTO attendance (
        student_id,
        full_name,
        class_name,
        subject,
        date,
        status,
        teacher_telegram_id,
        created_at
    )
    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
    """, (
        student_id,
        full_name,
        class_name,
        subject_name,
        today,
        status,
        teacher_telegram_id,
        datetime.now().isoformat()
    ))

    conn.commit()
    conn.close()


def attendance_by_date(target_date):
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT *
    FROM attendance
    WHERE date = ?
    ORDER BY class_name, subject, full_name
    """, (target_date,))

    rows = c.fetchall()
    conn.close()
    return rows


def stats():
    conn = db()
    c = conn.cursor()

    queries = {
        "students": "SELECT COUNT(*) AS count FROM students",
        "teachers": "SELECT COUNT(*) AS count FROM teachers",
        "subjects": "SELECT COUNT(*) AS count FROM subjects",
        "present": "SELECT COUNT(*) AS count FROM attendance WHERE status='present'",
        "absent": "SELECT COUNT(*) AS count FROM attendance WHERE status='absent'",
        "late": "SELECT COUNT(*) AS count FROM attendance WHERE status='late'",
    }

    result = {}

    for key, query in queries.items():
        c.execute(query)
        result[key] = c.fetchone()["count"]

    conn.close()
    return result


def top_absent(limit=10):
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT student_id, full_name, COUNT(*) AS total
    FROM attendance
    WHERE status = 'absent'
    GROUP BY student_id, full_name
    ORDER BY total DESC
    LIMIT ?
    """, (limit,))

    rows = c.fetchall()
    conn.close()
    return rows


def subject_stats():
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT
        subject,
        SUM(CASE WHEN status = 'present' THEN 1 ELSE 0 END) AS present,
        SUM(CASE WHEN status = 'absent' THEN 1 ELSE 0 END) AS absent,
        SUM(CASE WHEN status = 'late' THEN 1 ELSE 0 END) AS late
    FROM attendance
    GROUP BY subject
    ORDER BY absent DESC
    """)

    rows = c.fetchall()
    conn.close()
    return rows


def monthly(year_month):
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT
        SUM(CASE WHEN status = 'present' THEN 1 ELSE 0 END) AS present,
        SUM(CASE WHEN status = 'absent' THEN 1 ELSE 0 END) AS absent,
        SUM(CASE WHEN status = 'late' THEN 1 ELSE 0 END) AS late
    FROM attendance
    WHERE date LIKE ?
    """, (f"{year_month}%",))

    row = c.fetchone()
    conn.close()

    present = row["present"] or 0
    absent = row["absent"] or 0
    late = row["late"] or 0
    total = present + absent + late
    percent = round((present / total) * 100, 1) if total else 0

    return present, absent, late, total, percent


def student_summary(student_id):
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT
        SUM(CASE WHEN status = 'present' THEN 1 ELSE 0 END) AS present,
        SUM(CASE WHEN status = 'absent' THEN 1 ELSE 0 END) AS absent,
        SUM(CASE WHEN status = 'late' THEN 1 ELSE 0 END) AS late,
        COUNT(*) AS total
    FROM attendance
    WHERE student_id = ?
    """, (student_id,))

    summary = c.fetchone()

    c.execute("""
    SELECT date, subject, status
    FROM attendance
    WHERE student_id = ?
    ORDER BY date DESC, id DESC
    LIMIT 8
    """, (student_id,))

    records = c.fetchall()
    conn.close()

    present = summary["present"] or 0
    absent = summary["absent"] or 0
    late = summary["late"] or 0
    total = summary["total"] or 0
    percent = round((present / total) * 100, 1) if total else 0

    return present, absent, late, total, percent, records


def consecutive_absent(student_id, days=3):
    dates = [str(date.today() - timedelta(days=i)) for i in range(days)]
    placeholders = ",".join(["?"] * len(dates))

    conn = db()
    c = conn.cursor()

    c.execute(f"""
    SELECT COUNT(DISTINCT date) AS count
    FROM attendance
    WHERE student_id = ?
      AND status = 'absent'
      AND date IN ({placeholders})
    """, [student_id] + dates)

    count = c.fetchone()["count"]
    conn.close()

    return count >= days


def parent_ids():
    conn = db()
    c = conn.cursor()

    c.execute("""
    SELECT DISTINCT parent_telegram_id
    FROM students
    WHERE parent_telegram_id IS NOT NULL
    """)

    rows = c.fetchall()
    conn.close()

    return [row["parent_telegram_id"] for row in rows]


# ================= KEYBOARDS =================

def button(text, data):
    return InlineKeyboardButton(text=text, callback_data=data)


def keyboard(rows):
    return InlineKeyboardMarkup(inline_keyboard=rows)


def login_keyboard():
    return keyboard([
        [button("👨‍💼 Admin", "login:admin")],
        [button("👨‍🏫 O‘qituvchi", "login:teacher")],
        [button("👨‍👩‍👧 Ota-ona", "login:parent")],
    ])


def admin_keyboard():
    return keyboard([
        [button("👥 O‘quvchilar", "students")],
        [button("👨‍🏫 Teacherlar", "teachers")],
        [button("📘 Fanlar", "subjects")],
        [button("📊 Hisobotlar", "reports")],
        [button("📈 Analytics", "analytics")],
        [button("📤 Excel export", "excel")],
        [button("📄 PDF report", "pdf")],
        [button("📢 Broadcast", "broadcast")],
        [button("💾 Backup", "backup")],
    ])


def teacher_keyboard():
    return keyboard([
        [button("📚 Davomat olish", "attendance")],
        [button("➕ O‘quvchi qo‘shish", "add_student")],
        [button("👥 O‘quvchilar", "students")],
        [button("📘 Fanlar", "subjects")],
        [button("📊 Hisobotlar", "reports")],
        [button("📈 Analytics", "analytics")],
        [button("📤 Excel export", "excel")],
        [button("📄 PDF report", "pdf")],
        [button("🔍 O‘quvchi qidirish", "search")],
    ])


def parent_keyboard():
    return keyboard([
        [button("👶 Farzandlarim", "children")],
    ])


def menu_keyboard(user_role):
    if user_role == "admin":
        return admin_keyboard()

    if user_role == "teacher":
        return teacher_keyboard()

    if user_role == "parent":
        return parent_keyboard()

    return login_keyboard()


def status_text(status):
    return {
        "present": "✅ Keldi",
        "absent": "❌ Kelmadi",
        "late": "⏰ Kechikdi",
    }.get(status, status)


def status_icon(status):
    return {
        "present": "✅",
        "absent": "❌",
        "late": "⏰",
    }.get(status, "➖")


# ================= BOT =================

bot = Bot(BOT_TOKEN) if BOT_TOKEN else None
dp = Dispatcher()

STATE = {}
SESSION = {}
STATUS_ORDER = ["present", "absent", "late"]


def set_state(user_id, action, data=None):
    STATE[user_id] = {
        "action": action,
        "data": data or {}
    }


def pop_state(user_id):
    STATE.pop(user_id, None)


async def send_menu(message, user_role=None, text=None):
    user_role = user_role or role(message.from_user.id)

    await message.answer(
        text or ("Menyu:" if user_role else "🔐 Rol tanlang:"),
        reply_markup=menu_keyboard(user_role)
    )


async def notify_admins(text):
    if not bot:
        return

    for admin_id in ids_by_role("admin"):
        try:
            await bot.send_message(admin_id, text)
        except Exception as error:
            print("admin notify:", error)


async def notify_absence(student_id, full_name, class_name, subject_name, today):
    if not bot:
        return

    current_student = student(student_id)

    if current_student and current_student["parent_telegram_id"]:
        try:
            await bot.send_message(
                current_student["parent_telegram_id"],
                (
                    "❌ Davomat bildirishnomasi\n\n"
                    f"Farzandingiz {full_name} bugun darsga kelmadi.\n\n"
                    f"🏫 Sinf: {class_name}\n"
                    f"📘 Fan: {subject_name}\n"
                    f"📅 Sana: {today}"
                )
            )
        except Exception as error:
            print("parent notify:", error)

    await notify_admins(
        (
            "❌ Kelmadi\n\n"
            f"👤 {full_name}\n"
            f"🏫 {class_name}\n"
            f"📘 {subject_name}\n"
            f"📅 {today}"
        )
    )

    if consecutive_absent(student_id, 3):
        alert = (
            "⚠️ 3 kunlik ogohlantirish\n\n"
            f"👤 {full_name}\n"
            f"🏫 {class_name}\n"
            "❌ Ketma-ket kelmagan holat bor."
        )

        await notify_admins(alert)

        if current_student and current_student["parent_telegram_id"]:
            try:
                await bot.send_message(current_student["parent_telegram_id"], alert)
            except Exception:
                pass


def profile_text(student_id):
    current_student = student(student_id)

    if not current_student:
        return "❌ O‘quvchi topilmadi."

    present, absent, late, total, percent, records = student_summary(student_id)

    text = (
        "👤 O‘quvchi profili\n\n"
        f"🆔 ID: {current_student['id']}\n"
        f"👤 Ism: {current_student['full_name']}\n"
        f"🏫 Sinf: {current_student['class_name']}\n"
        f"📞 Telefon: {current_student['parent_phone'] or '-'}\n"
        f"🔐 Parent kodi: {current_student['parent_code']}\n\n"
        "📊 Davomat\n"
        f"✅ Keldi: {present}\n"
        f"❌ Kelmadi: {absent}\n"
        f"⏰ Kechikdi: {late}\n"
        f"📈 Foiz: {percent}%\n"
    )

    if records:
        text += "\nOxirgi yozuvlar:\n"
        for record in records:
            text += f"{record['date']} | {record['subject']} | {status_text(record['status'])}\n"

    return text


@dp.message(CommandStart())
async def start(message: Message):
    current_role = role(message.from_user.id)

    if current_role:
        await send_menu(message, current_role)
    else:
        await message.answer(
            "🔐 Botga kirish uchun rolingizni tanlang:",
            reply_markup=login_keyboard()
        )


@dp.callback_query(F.data.startswith("login:"))
async def login_choose(callback: CallbackQuery):
    selected_role = callback.data.split(":")[1]

    set_state(
        callback.from_user.id,
        "login",
        {"role": selected_role}
    )

    if selected_role == "admin":
        await callback.message.answer("👨‍💼 Admin parolini kiriting:")

    elif selected_role == "teacher":
        await callback.message.answer("👨‍🏫 O‘qituvchi paroli yoki TCH-kodni kiriting:")

    else:
        await callback.message.answer("👨‍👩‍👧 Ota-ona PAR-kodni kiriting:")

    await callback.answer()


@dp.message()
async def text_handler(message: Message):
    user_id = message.from_user.id
    current_state = STATE.get(user_id)
    message_text = (message.text or "").strip()

    if not current_state:
        await message.answer("Buyruq tushunarsiz. /start ni bosing.")
        return

    action = current_state["action"]
    data = current_state["data"]

    if action == "login":
        selected_role = data["role"]

        if selected_role == "admin":
            if message_text != ADMIN_CODE:
                await message.answer("❌ Admin paroli xato.")
                return

            set_role(user_id, "admin", message.from_user.full_name)
            pop_state(user_id)
            await send_menu(message, "admin", "✅ Admin sifatida kirdingiz.")
            return

        if selected_role == "teacher":
            if message_text == TEACHER_CODE:
                set_role(user_id, "teacher", message.from_user.full_name)
                pop_state(user_id)
                await send_menu(message, "teacher", "✅ Teacher sifatida kirdingiz.")
                return

            if message_text.upper().startswith("TCH-") and link_teacher(
                message_text,
                user_id,
                message.from_user.full_name
            ):
                pop_state(user_id)
                await send_menu(message, "teacher", "✅ Teacher kodingiz tasdiqlandi.")
                return

            await message.answer("❌ Teacher paroli/kodi xato.")
            return

        if selected_role == "parent":
            if link_parent(message_text, user_id, message.from_user.full_name):
                pop_state(user_id)
                await send_menu(message, "parent", "✅ Parent sifatida kirdingiz.")
                return

            await message.answer("❌ Parent kodi xato.")
            return

    if action == "add_student":
        if role(user_id) != "teacher":
            pop_state(user_id)
            await message.answer("❌ Bu bo‘lim faqat teacher uchun.")
            return

        if message_text.count(",") < 2:
            await message.answer("Format: Ali Karimov, 10-A, +998901234567")
            return

        full_name, class_name, phone = [
            item.strip()
            for item in message_text.split(",", 2)
        ]

        parent_code = add_student(full_name, class_name, phone)
        pop_state(user_id)

        await message.answer(
            (
                "✅ O‘quvchi qo‘shildi\n\n"
                f"👤 {full_name}\n"
                f"🏫 {class_name}\n"
                f"📞 {phone}\n"
                f"🔐 Parent kodi: {parent_code}"
            ),
            reply_markup=teacher_keyboard()
        )
        return

    if action == "search":
        rows = search_students(message_text)
        pop_state(user_id)
        current_role = role(user_id)

        if not rows:
            await message.answer("❌ Topilmadi.", reply_markup=menu_keyboard(current_role))
            return

        text = "🔍 Natijalar:\n\n"
        buttons = []

        for row in rows:
            text += (
                f"🆔 {row['id']} | "
                f"{row['full_name']} | "
                f"{row['class_name']} | "
                f"{row['parent_code']}\n"
            )

            buttons.append([
                button(f"👤 {row['full_name']}", f"profile:{row['id']}")
            ])

        buttons.append([button("⬅️ Menyu", "menu")])

        await message.answer(text, reply_markup=keyboard(buttons))
        return

    if action == "subject":
        add_subject(message_text)
        pop_state(user_id)

        await message.answer(
            f"✅ Fan qo‘shildi: {message_text}",
            reply_markup=menu_keyboard(role(user_id))
        )
        return

    if action == "teacher":
        if role(user_id) != "admin":
            pop_state(user_id)
            await message.answer("❌ Bu bo‘lim faqat admin uchun.")
            return

        parts = [item.strip() for item in message_text.split(",") if item.strip()]
        teacher_name = parts[0]
        class_names = ",".join(parts[1:]) if len(parts) > 1 else ""

        teacher_code = add_teacher(teacher_name, class_names)
        pop_state(user_id)

        await message.answer(
            (
                "✅ Teacher qo‘shildi\n"
                f"👨‍🏫 {teacher_name}\n"
                f"🔐 Kod: {teacher_code}\n"
                f"🏫 Sinflar: {', '.join(parts[1:]) or 'hammasi'}"
            ),
            reply_markup=admin_keyboard()
        )
        return

    if action == "broadcast":
        if role(user_id) != "admin":
            pop_state(user_id)
            await message.answer("❌ Bu bo‘lim faqat admin uchun.")
            return

        success = 0

        for parent_id in parent_ids():
            try:
                await bot.send_message(
                    parent_id,
                    f"📢 Maktab xabari\n\n{message_text}"
                )
                success += 1
            except Exception as error:
                print("broadcast:", error)

        pop_state(user_id)

        await message.answer(
            f"✅ Xabar yuborildi: {success}",
            reply_markup=admin_keyboard()
        )
        return


@dp.callback_query(F.data == "menu")
async def menu(callback: CallbackQuery):
    await callback.message.answer(
        "Menyu:",
        reply_markup=menu_keyboard(role(callback.from_user.id))
    )
    await callback.answer()


@dp.callback_query(F.data == "add_student")
async def callback_add_student(callback: CallbackQuery):
    if role(callback.from_user.id) != "teacher":
        await callback.answer("Faqat teacher.", show_alert=True)
        return

    set_state(callback.from_user.id, "add_student")
    await callback.message.answer("O‘quvchi: Ali Karimov, 10-A, +998901234567")
    await callback.answer()


@dp.callback_query(F.data == "students")
async def callback_students(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    rows = students()
    text = "👥 O‘quvchilar:\n\n"
    buttons = []

    for row in rows:
        text += (
            f"🆔 {row['id']} | "
            f"{row['full_name']} | "
            f"{row['class_name']} | "
            f"Kod: {row['parent_code']}\n"
        )

        buttons.append([
            button(f"👤 {row['full_name']}", f"profile:{row['id']}")
        ])

    buttons.append([button("🗑 O‘chirish", "delete_menu")])
    buttons.append([button("⬅️ Menyu", "menu")])

    await callback.message.answer(
        text if rows else "O‘quvchi yo‘q.",
        reply_markup=keyboard(buttons)
    )

    await callback.answer()


@dp.callback_query(F.data == "search")
async def callback_search(callback: CallbackQuery):
    set_state(callback.from_user.id, "search")
    await callback.message.answer("🔍 Ism, sinf, telefon yoki kod kiriting:")
    await callback.answer()


@dp.callback_query(F.data.startswith("profile:"))
async def callback_profile(callback: CallbackQuery):
    student_id = int(callback.data.split(":")[1])
    current_role = role(callback.from_user.id)

    if current_role == "parent":
        child_ids = [child["id"] for child in parent_children(callback.from_user.id)]

        if student_id not in child_ids:
            await callback.answer("Bu profil sizniki emas.", show_alert=True)
            return

    await callback.message.answer(
        profile_text(student_id),
        reply_markup=menu_keyboard(current_role)
    )

    await callback.answer()


@dp.callback_query(F.data == "children")
async def callback_children(callback: CallbackQuery):
    rows = parent_children(callback.from_user.id)
    buttons = []

    for row in rows:
        buttons.append([
            button(
                f"👤 {row['full_name']} — {row['class_name']}",
                f"profile:{row['id']}"
            )
        ])

    buttons.append([button("⬅️ Menyu", "menu")])

    await callback.message.answer(
        "Farzandlaringiz:" if rows else "Farzand topilmadi.",
        reply_markup=keyboard(buttons)
    )

    await callback.answer()


@dp.callback_query(F.data == "delete_menu")
async def callback_delete_menu(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    buttons = []

    for row in students():
        buttons.append([
            button(
                f"🗑 {row['full_name']} — {row['class_name']}",
                f"delete:{row['id']}"
            )
        ])

    buttons.append([button("⬅️ Menyu", "menu")])

    await callback.message.answer(
        "O‘chirish:",
        reply_markup=keyboard(buttons)
    )

    await callback.answer()


@dp.callback_query(F.data.startswith("delete:"))
async def callback_delete(callback: CallbackQuery):
    student_id = int(callback.data.split(":")[1])
    delete_student(student_id)

    await callback.message.answer(
        "✅ O‘quvchi o‘chirildi.",
        reply_markup=menu_keyboard(role(callback.from_user.id))
    )

    await callback.answer()


@dp.callback_query(F.data == "subjects")
async def callback_subjects(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    subject_list = subjects()

    text = "📘 Fanlar:\n\n"

    if subject_list:
        text += "\n".join(f"• {item}" for item in subject_list)
    else:
        text += "Fan yo‘q."

    await callback.message.answer(
        text,
        reply_markup=keyboard([
            [button("➕ Fan qo‘shish", "add_subject")],
            [button("⬅️ Menyu", "menu")]
        ])
    )

    await callback.answer()


@dp.callback_query(F.data == "add_subject")
async def callback_add_subject(callback: CallbackQuery):
    set_state(callback.from_user.id, "subject")
    await callback.message.answer("Fan nomini kiriting:")
    await callback.answer()


@dp.callback_query(F.data == "teachers")
async def callback_teachers(callback: CallbackQuery):
    if role(callback.from_user.id) != "admin":
        await callback.answer("Faqat admin.", show_alert=True)
        return

    text = "👨‍🏫 Teacherlar:\n\n"

    for row in teachers():
        text += (
            f"🆔 {row['id']} | "
            f"{row['full_name']} | "
            f"{row['teacher_code']} | "
            f"{'✅ Ulangan' if row['telegram_id'] else '⏳ Ulanmagan'}\n"
        )

    await callback.message.answer(
        text,
        reply_markup=keyboard([
            [button("➕ Teacher qo‘shish", "add_teacher")],
            [button("⬅️ Menyu", "menu")]
        ])
    )

    await callback.answer()


@dp.callback_query(F.data == "add_teacher")
async def callback_add_teacher(callback: CallbackQuery):
    set_state(callback.from_user.id, "teacher")
    await callback.message.answer(
        "Format:\n"
        "Ali Ustoz, 9-A, 10-B\n\n"
        "Agar hamma sinflarni ko‘rsin desangiz:\n"
        "Ali Ustoz"
    )
    await callback.answer()


# ================= ATTENDANCE =================

@dp.callback_query(F.data == "attendance")
async def callback_attendance(callback: CallbackQuery):
    if role(callback.from_user.id) != "teacher":
        await callback.answer("Faqat teacher.", show_alert=True)
        return

    class_list = classes_for_teacher(callback.from_user.id)

    if not class_list:
        await callback.message.answer("Hali sinf/o‘quvchi yo‘q.")
        await callback.answer()
        return

    buttons = [
        [button(class_name, f"attendance_class:{class_name}")]
        for class_name in class_list
    ]

    buttons.append([button("⬅️ Menyu", "menu")])

    await callback.message.answer(
        "🏫 Sinfni tanlang:",
        reply_markup=keyboard(buttons)
    )

    await callback.answer()


@dp.callback_query(F.data.startswith("attendance_class:"))
async def callback_attendance_class(callback: CallbackQuery):
    class_name = callback.data.split(":", 1)[1]

    SESSION[callback.from_user.id] = {
        "class": class_name,
        "subject": None,
        "students": {}
    }

    subject_list = subjects()

    if not subject_list:
        await callback.message.answer("Avval fan qo‘shing.")
        await callback.answer()
        return

    buttons = [
        [button(subject_name, f"attendance_subject:{subject_name}")]
        for subject_name in subject_list
    ]

    buttons.append([button("⬅️ Menyu", "menu")])

    await callback.message.answer(
        f"📘 Fanni tanlang\n\n🏫 {class_name}",
        reply_markup=keyboard(buttons)
    )

    await callback.answer()


async def render_attendance(callback: CallbackQuery, edit=True):
    current_session = SESSION[callback.from_user.id]
    rows = students_by_class(current_session["class"])

    text = (
        "📚 Davomat\n\n"
        f"🏫 {current_session['class']}\n"
        f"📘 {current_session['subject']}\n"
        f"📅 {date.today()}\n\n"
    )

    buttons = []

    for row in rows:
        current_status = current_session["students"].get(row["id"], "present")

        text += f"{status_icon(current_status)} {row['full_name']}\n"

        buttons.append([
            button(
                f"{status_icon(current_status)} {row['full_name']}",
                f"attendance_toggle:{row['id']}"
            )
        ])

    buttons.append([button("💾 Saqlash", "attendance_save")])
    buttons.append([button("⬅️ Menyu", "menu")])

    if edit:
        await callback.message.edit_text(
            text,
            reply_markup=keyboard(buttons)
        )
    else:
        await callback.message.answer(
            text,
            reply_markup=keyboard(buttons)
        )


@dp.callback_query(F.data.startswith("attendance_subject:"))
async def callback_attendance_subject(callback: CallbackQuery):
    current_session = SESSION.get(callback.from_user.id)

    if not current_session:
        await callback.answer("Avval sinf tanlang.", show_alert=True)
        return

    current_session["subject"] = callback.data.split(":", 1)[1]
    current_session["students"] = {
        row["id"]: "present"
        for row in students_by_class(current_session["class"])
    }

    await render_attendance(callback, edit=False)
    await callback.answer()


@dp.callback_query(F.data.startswith("attendance_toggle:"))
async def callback_attendance_toggle(callback: CallbackQuery):
    current_session = SESSION.get(callback.from_user.id)

    if not current_session:
        await callback.answer("Davomat sessiyasi yo‘q.", show_alert=True)
        return

    student_id = int(callback.data.split(":")[1])
    current_status = current_session["students"].get(student_id, "present")

    next_status = STATUS_ORDER[
        (STATUS_ORDER.index(current_status) + 1) % len(STATUS_ORDER)
    ]

    current_session["students"][student_id] = next_status

    await render_attendance(callback, edit=True)
    await callback.answer()


@dp.callback_query(F.data == "attendance_save")
async def callback_attendance_save(callback: CallbackQuery):
    current_session = SESSION.get(callback.from_user.id)

    if not current_session:
        await callback.answer("Sessiya yo‘q.", show_alert=True)
        return

    today = str(date.today())

    counts = {
        "present": 0,
        "absent": 0,
        "late": 0
    }

    absent_students = []

    for row in students_by_class(current_session["class"]):
        current_status = current_session["students"].get(row["id"], "present")
        counts[current_status] += 1

        save_attendance(
            row["id"],
            row["full_name"],
            current_session["class"],
            current_session["subject"],
            current_status,
            callback.from_user.id
        )

        if current_status == "absent":
            absent_students.append(row["full_name"])
            await notify_absence(
                row["id"],
                row["full_name"],
                current_session["class"],
                current_session["subject"],
                today
            )

    SESSION.pop(callback.from_user.id, None)

    text = (
        "✅ Davomat saqlandi\n\n"
        f"🏫 {current_session['class']}\n"
        f"📘 {current_session['subject']}\n"
        f"✅ Keldi: {counts['present']}\n"
        f"❌ Kelmadi: {counts['absent']}\n"
        f"⏰ Kechikdi: {counts['late']}"
    )

    if absent_students:
        text += "\n\nKelmaganlar:\n" + "\n".join(absent_students)

    await callback.message.answer(text, reply_markup=teacher_keyboard())
    await callback.answer()


# ================= REPORTS =================

@dp.callback_query(F.data == "reports")
async def callback_reports(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    current_stats = stats()
    current_month = str(date.today())[:7]
    present, absent, late, total, percent = monthly(current_month)

    text = (
        "📊 Hisobotlar\n\n"
        f"👥 O‘quvchilar: {current_stats['students']}\n"
        f"👨‍🏫 Teacherlar: {current_stats['teachers']}\n"
        f"📘 Fanlar: {current_stats['subjects']}\n\n"
        f"✅ Keldi: {current_stats['present']}\n"
        f"❌ Kelmadi: {current_stats['absent']}\n"
        f"⏰ Kechikdi: {current_stats['late']}\n\n"
        f"📅 {current_month}: ✅ {present} ❌ {absent} ⏰ {late} 📈 {percent}%\n\n"
        "📈 TOP kelmaganlar:\n"
    )

    absent_rows = top_absent()

    if absent_rows:
        for index, row in enumerate(absent_rows, 1):
            text += f"{index}. {row['full_name']} — {row['total']}\n"
    else:
        text += "Ma’lumot yo‘q.\n"

    await callback.message.answer(
        text,
        reply_markup=keyboard([
            [button("📤 Excel", "excel")],
            [button("📄 PDF", "pdf")],
            [button("⬅️ Menyu", "menu")]
        ])
    )

    await callback.answer()


@dp.callback_query(F.data == "analytics")
async def callback_analytics(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    text = "📈 Fanlar statistikasi:\n\n"

    rows = subject_stats()

    if rows:
        for row in rows:
            text += (
                f"• {row['subject']}: "
                f"✅ {row['present'] or 0} | "
                f"❌ {row['absent'] or 0} | "
                f"⏰ {row['late'] or 0}\n"
            )
    else:
        text += "Ma’lumot yo‘q."

    await callback.message.answer(
        text,
        reply_markup=menu_keyboard(role(callback.from_user.id))
    )

    await callback.answer()


@dp.callback_query(F.data == "excel")
async def callback_excel(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    rows = attendance_by_date(str(date.today()))

    if not rows:
        await callback.message.answer("Bugungi davomat yo‘q.")
        await callback.answer()
        return

    path = f"exports/excel/attendance_{date.today()}.xlsx"

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Davomat"

    sheet.append(["ID", "F.I.SH", "Sinf", "Fan", "Status", "Sana"])

    for row in rows:
        sheet.append([
            row["id"],
            row["full_name"],
            row["class_name"],
            row["subject"],
            status_text(row["status"]),
            row["date"]
        ])

    workbook.save(path)

    await callback.message.answer_document(FSInputFile(path))
    await callback.answer()


@dp.callback_query(F.data == "pdf")
async def callback_pdf(callback: CallbackQuery):
    if role(callback.from_user.id) not in ["admin", "teacher"]:
        await callback.answer("Ruxsat yo‘q.", show_alert=True)
        return

    rows = attendance_by_date(str(date.today()))

    if not rows:
        await callback.message.answer("Bugungi davomat yo‘q.")
        await callback.answer()
        return

    path = f"exports/pdf/attendance_{date.today()}.pdf"

    pdf = canvas.Canvas(path, pagesize=A4)
    y = 800

    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(50, y, "School ERP - Davomat hisoboti")
    y -= 30

    pdf.setFont("Helvetica", 10)

    for row in rows:
        if y < 50:
            pdf.showPage()
            y = 800
            pdf.setFont("Helvetica", 10)

        pdf.drawString(
            50,
            y,
            (
                f"ID: {row['id']} | "
                f"{row['full_name']} | "
                f"{row['class_name']} | "
                f"{row['subject']} | "
                f"{status_text(row['status'])} | "
                f"{row['date']}"
            )
        )
        y -= 15

    pdf.save()

    await callback.message.answer_document(FSInputFile(path))
    await callback.answer()


@dp.callback_query(F.data == "broadcast")
async def callback_broadcast(callback: CallbackQuery):
    if role(callback.from_user.id) != "admin":
        await callback.answer("Faqat admin.", show_alert=True)
        return

    set_state(callback.from_user.id, "broadcast")
    await callback.message.answer("Broadcast matnini kiriting:")
    await callback.answer()


@dp.callback_query(F.data == "backup")
async def callback_backup(callback: CallbackQuery):
    if role(callback.from_user.id) != "admin":
        await callback.answer("Faqat admin.", show_alert=True)
        return

    path = f"backups/davomat_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    shutil.copy(DB_NAME, path)

    await callback.message.answer_document(FSInputFile(path))
    await callback.answer()


# ================= WEB DASHBOARD =================

def esc(value):
    return html.escape(str(value if value is not None else ""))


def badge(status):
    status_label = status_text(status)

    css_class = {
        "present": "present",
        "absent": "absent",
        "late": "late",
    }.get(status, "")

    return f'<span class="badge {css_class}">{esc(status_label)}</span>'


def layout(title, body):
    return f"""
    <!doctype html>
    <html lang="uz">
    <head>
        <meta charset="utf-8">
        <meta name="viewport" content="width=device-width, initial-scale=1">
        <title>{esc(title)}</title>
        <style>
            * {{ box-sizing: border-box; }}
            body {{
                margin: 0;
                font-family: Arial, sans-serif;
                background: #f4f6f9;
                color: #111827;
            }}
            .sidebar {{
                width: 260px;
                height: 100vh;
                position: fixed;
                top: 0;
                left: 0;
                background: #111827;
                color: white;
                padding: 24px 18px;
            }}
            .sidebar h2 {{
                margin: 0 0 28px;
                font-size: 22px;
            }}
            .sidebar a {{
                display: block;
                padding: 12px 14px;
                margin-bottom: 8px;
                color: #d1d5db;
                text-decoration: none;
                border-radius: 10px;
                font-size: 15px;
            }}
            .sidebar a:hover {{
                background: #374151;
                color: white;
            }}
            .main {{
                margin-left: 260px;
                padding: 30px;
            }}
            h1 {{
                margin-top: 0;
            }}
            .cards {{
                display: grid;
                grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
                gap: 18px;
                margin-bottom: 28px;
            }}
            .card {{
                background: white;
                padding: 22px;
                border-radius: 16px;
                box-shadow: 0 4px 20px rgba(0,0,0,0.06);
            }}
            .card span {{
                color: #6b7280;
                font-size: 14px;
            }}
            .card strong {{
                display: block;
                margin-top: 8px;
                font-size: 30px;
            }}
            .table-wrap {{
                background: white;
                border-radius: 16px;
                overflow: hidden;
                box-shadow: 0 4px 20px rgba(0,0,0,0.06);
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
            }}
            th, td {{
                padding: 13px 14px;
                border-bottom: 1px solid #e5e7eb;
                text-align: left;
                font-size: 14px;
            }}
            th {{
                background: #111827;
                color: white;
            }}
            tr:hover td {{
                background: #f9fafb;
            }}
            .badge {{
                display: inline-block;
                padding: 5px 10px;
                border-radius: 999px;
                font-weight: 700;
                font-size: 12px;
            }}
            .present {{ background: #dcfce7; color: #166534; }}
            .absent {{ background: #fee2e2; color: #991b1b; }}
            .late {{ background: #fef3c7; color: #92400e; }}
            @media (max-width: 800px) {{
                .sidebar {{
                    position: static;
                    width: 100%;
                    height: auto;
                }}
                .main {{
                    margin-left: 0;
                    padding: 18px;
                }}
                table {{
                    min-width: 760px;
                }}
                .table-wrap {{
                    overflow-x: auto;
                }}
            }}
        </style>
    </head>
    <body>
        <aside class="sidebar">
            <h2>School ERP</h2>
            <a href="/">📊 Dashboard</a>
            <a href="/students">👥 O‘quvchilar</a>
            <a href="/teachers">👨‍🏫 Teacherlar</a>
            <a href="/subjects">📘 Fanlar</a>
            <a href="/attendance">📚 Davomat</a>
            <a href="/reports">📈 Reports</a>
        </aside>
        <main class="main">
            {body}
        </main>
    </body>
    </html>
    """


def table(headers, rows):
    header_html = "".join(f"<th>{esc(header)}</th>" for header in headers)

    row_html = ""

    for row in rows:
        row_html += "<tr>"
        for cell in row:
            row_html += f"<td>{cell}</td>"
        row_html += "</tr>"

    return f"""
    <div class="table-wrap">
        <table>
            <thead>
                <tr>{header_html}</tr>
            </thead>
            <tbody>
                {row_html}
            </tbody>
        </table>
    </div>
    """


@web_app.get("/", response_class=HTMLResponse)
async def web_dashboard():
    current_stats = stats()
    body = f"""
    <h1>📊 Dashboard</h1>
    <div class="cards">
        <div class="card"><span>O‘quvchilar</span><strong>{current_stats['students']}</strong></div>
        <div class="card"><span>Teacherlar</span><strong>{current_stats['teachers']}</strong></div>
        <div class="card"><span>Fanlar</span><strong>{current_stats['subjects']}</strong></div>
        <div class="card"><span>Keldi</span><strong>{current_stats['present']}</strong></div>
        <div class="card"><span>Kelmadi</span><strong>{current_stats['absent']}</strong></div>
        <div class="card"><span>Kechikdi</span><strong>{current_stats['late']}</strong></div>
    </div>
    """
    return layout("Dashboard", body)


@web_app.get("/students", response_class=HTMLResponse)
async def web_students():
    rows = []

    for row in students():
        rows.append([
            esc(row["id"]),
            esc(row["full_name"]),
            esc(row["class_name"]),
            esc(row["parent_phone"]),
            esc(row["parent_code"]),
            esc(row["parent_telegram_id"] or "")
        ])

    body = "<h1>👥 O‘quvchilar</h1>" + table(
        ["ID", "Ism familiya", "Sinf", "Telefon", "Parent kod", "Parent Telegram ID"],
        rows
    )

    return layout("O‘quvchilar", body)


@web_app.get("/teachers", response_class=HTMLResponse)
async def web_teachers():
    rows = []

    for row in teachers():
        rows.append([
            esc(row["id"]),
            esc(row["full_name"]),
            esc(row["telegram_id"] or ""),
            esc(row["teacher_code"]),
            "✅ Ulangan" if row["telegram_id"] else "⏳ Ulanmagan"
        ])

    body = "<h1>👨‍🏫 Teacherlar</h1>" + table(
        ["ID", "Ism familiya", "Telegram ID", "Teacher kod", "Status"],
        rows
    )

    return layout("Teacherlar", body)


@web_app.get("/subjects", response_class=HTMLResponse)
async def web_subjects():
    subject_list = subjects()
    rows = [
        [esc(index), esc(name)]
        for index, name in enumerate(subject_list, start=1)
    ]

    body = "<h1>📘 Fanlar</h1>" + table(
        ["ID", "Fan nomi"],
        rows
    )

    return layout("Fanlar", body)


@web_app.get("/attendance", response_class=HTMLResponse)
async def web_attendance():
    rows = []

    for row in attendance_by_date(str(date.today())):
        rows.append([
            esc(row["id"]),
            esc(row["full_name"]),
            esc(row["class_name"]),
            esc(row["subject"]),
            badge(row["status"]),
            esc(row["date"]),
            esc(row["teacher_telegram_id"] or "")
        ])

    body = "<h1>📚 Bugungi davomat</h1>" + table(
        ["ID", "O‘quvchi", "Sinf", "Fan", "Status", "Sana", "Teacher Telegram ID"],
        rows
    )

    return layout("Davomat", body)


@web_app.get("/reports", response_class=HTMLResponse)
async def web_reports():
    rows = []

    for row in top_absent():
        rows.append([
            esc(row["student_id"]),
            esc(row["full_name"]),
            esc(row["total"])
        ])

    body = "<h1>📈 TOP kelmaganlar</h1>" + table(
        ["Student ID", "O‘quvchi", "Kelmadi soni"],
        rows
    )

    return layout("Reports", body)


def run_web():
    uvicorn.run(
        web_app,
        host=WEB_HOST,
        port=WEB_PORT,
        log_level="warning"
    )


async def main():
    if not BOT_TOKEN:
        raise RuntimeError("BOT_TOKEN .env da yo‘q")

    init_db()

    threading.Thread(
        target=run_web,
        daemon=True
    ).start()

    print("Bot ishga tushdi.")
    print(f"Web local: http://127.0.0.1:{WEB_PORT}")
    print(f"Web host:  http://{WEB_HOST}:{WEB_PORT}")

    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
